from flask import Flask, request, jsonify, send_from_directory
from flask_cors import CORS
from werkzeug.utils import secure_filename
from dotenv import load_dotenv
from pathlib import Path
from datetime import datetime, timezone
import hashlib
import hmac
import threading
import time
import json
import mimetypes
import os
import secrets
import uuid
import urllib.request
import urllib.error

from google import genai
from google.genai import types

import mysql.connector
from mysql.connector import Error

BASE = Path(__file__).resolve().parent
load_dotenv(BASE / ".env")
DATA_DIR = BASE / "data"
UPLOAD_DIR = BASE / "uploads"
DB_FILE = DATA_DIR / "db.json"  # legacy one-time migration source
DATA_DIR.mkdir(parents=True, exist_ok=True)
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

app = Flask(__name__)
CORS(app)
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024

MYSQL_HOST = os.getenv("MYSQL_HOST", "127.0.0.1")
MYSQL_PORT = int(os.getenv("MYSQL_PORT", "3306"))
MYSQL_USER = os.getenv("MYSQL_USER", "root")
MYSQL_PASSWORD = os.getenv("MYSQL_PASSWORD", "")
MYSQL_DATABASE = os.getenv("MYSQL_DATABASE", "islearning")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "").strip()
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-2.5-flash").strip()
AI_PROVIDER_MODE = os.getenv("AI_PROVIDER_MODE", "hybrid").strip().lower()
OLLAMA_BASE_URL = os.getenv("OLLAMA_BASE_URL", "http://127.0.0.1:11434").rstrip("/")
OLLAMA_MODEL = os.getenv("OLLAMA_MODEL", "qwen2.5:3b").strip()
OLLAMA_TIMEOUT_SECONDS = int(os.getenv("OLLAMA_TIMEOUT_SECONDS", "240"))
# Real Gemini fallback chain. These are all current stable Flash models; the
# backend will only move to the next model when Google reports a temporary
# capacity/rate-limit problem or the configured model is unavailable.
GEMINI_FALLBACK_MODELS = [
    x.strip() for x in os.getenv(
        "GEMINI_MODEL_FALLBACKS",
        "gemini-2.5-flash-lite,gemini-2.0-flash"
    ).split(",") if x.strip()
]

DEFAULT_SETTINGS = {
    "darkMode": False,
    "offlineMode": False,
    "notificationsEnabled": True,
    "autoSave": True,
    "offlineAccess": True,
}


def mysql_connection(database=True):
    config = {
        "host": MYSQL_HOST,
        "port": MYSQL_PORT,
        "user": MYSQL_USER,
        "password": MYSQL_PASSWORD,
    }
    if database:
        config["database"] = MYSQL_DATABASE
    return mysql.connector.connect(**config)


def uid(prefix=""):
    return prefix + str(uuid.uuid4())


def now_label():
    if os.name == "nt":
        return datetime.now().strftime("%b %#d, %Y")
    return datetime.now().strftime("%b %-d, %Y")


def utc_now():
    return datetime.now(timezone.utc).replace(tzinfo=None)


def jload(value, default):
    try:
        return json.loads(value) if value else default
    except Exception:
        return default


def jdump(value):
    return json.dumps(value, ensure_ascii=False)


def password_hash(password):
    """Password storage for new/changed passwords. Legacy SHA-256 is still accepted."""
    iterations = 210_000
    salt = secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, iterations)
    return f"pbkdf2_sha256${iterations}${salt.hex()}${digest.hex()}"


def verify_password(password, stored):
    if not stored:
        return False
    if stored.startswith("pbkdf2_sha256$"):
        try:
            _, iters, salt_hex, digest_hex = stored.split("$", 3)
            actual = hashlib.pbkdf2_hmac(
                "sha256", password.encode(), bytes.fromhex(salt_hex), int(iters)
            ).hex()
            return hmac.compare_digest(actual, digest_hex)
        except Exception:
            return False
    # Backward compatibility with the original project.
    return hmac.compare_digest(
        stored, hashlib.sha256(password.encode()).hexdigest()
    )


def init_database():
    conn = mysql_connection(database=False)
    cur = conn.cursor()
    cur.execute(
        f"CREATE DATABASE IF NOT EXISTS `{MYSQL_DATABASE}` "
        "CHARACTER SET utf8mb4 COLLATE utf8mb4_unicode_ci"
    )
    cur.close()
    conn.close()

    conn = mysql_connection(database=True)
    cur = conn.cursor()

    # users keeps the original JSON columns temporarily so existing installations
    # can migrate safely. The application no longer depends on those columns.
    cur.execute("""
        CREATE TABLE IF NOT EXISTS users (
            id VARCHAR(100) PRIMARY KEY,
            student_id VARCHAR(100) NULL,
            email VARCHAR(255) NULL,
            name VARCHAR(255) NOT NULL DEFAULT '',
            first_name VARCHAR(255) NOT NULL DEFAULT '',
            middle_name VARCHAR(255) NOT NULL DEFAULT '',
            last_name VARCHAR(255) NOT NULL DEFAULT '',
            course VARCHAR(255) NOT NULL DEFAULT '',
            is_setup_complete TINYINT(1) NOT NULL DEFAULT 0,
            password_hash VARCHAR(255) NULL,
            is_guest TINYINT(1) NOT NULL DEFAULT 0,
            materials LONGTEXT NOT NULL,
            reviewers LONGTEXT NOT NULL,
            quiz_history LONGTEXT NOT NULL,
            notifications LONGTEXT NOT NULL,
            settings LONGTEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            updated_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
            INDEX idx_student_id (student_id),
            INDEX idx_email (email)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS material_folders (
            id VARCHAR(100) PRIMARY KEY,
            user_id VARCHAR(100) NOT NULL,
            name VARCHAR(255) NOT NULL,
            created_at DATETIME NOT NULL,
            CONSTRAINT fk_folder_user FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE,
            UNIQUE KEY uq_folder_name (user_id, name),
            INDEX idx_folder_user (user_id)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS materials (
            id VARCHAR(100) PRIMARY KEY,
            user_id VARCHAR(100) NOT NULL,
            title VARCHAR(255) NOT NULL,
            original_filename VARCHAR(255) NOT NULL,
            subject VARCHAR(255) NOT NULL DEFAULT 'General',
            topic VARCHAR(255) NOT NULL DEFAULT 'General Topic',
            week VARCHAR(50) NOT NULL DEFAULT 'Week 1',
            folder_id VARCHAR(100) NULL,
            file_type VARCHAR(20) NOT NULL,
            mime_type VARCHAR(150) NULL,
            stored_filename VARCHAR(500) NOT NULL,
            file_path VARCHAR(1000) NOT NULL,
            file_size_bytes BIGINT UNSIGNED NOT NULL DEFAULT 0,
            uploaded_at DATETIME NOT NULL,
            updated_at DATETIME NOT NULL,
            can_generate_ai TINYINT(1) NOT NULL DEFAULT 0,
            CONSTRAINT fk_material_user FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE,
            INDEX idx_material_user (user_id),
            INDEX idx_material_subject (user_id, subject),
            INDEX idx_material_uploaded (user_id, uploaded_at)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
    """)

    # Backward-compatible profile migrations for existing databases.
    for statement in (
        "ALTER TABLE users ADD COLUMN first_name VARCHAR(255) NOT NULL DEFAULT ''",
        "ALTER TABLE users ADD COLUMN middle_name VARCHAR(255) NOT NULL DEFAULT ''",
        "ALTER TABLE users ADD COLUMN last_name VARCHAR(255) NOT NULL DEFAULT ''",
    ):
        try:
            cur.execute(statement)
        except Error:
            pass

    cur.execute("""
        CREATE TABLE IF NOT EXISTS reviewers (
            id VARCHAR(100) PRIMARY KEY,
            user_id VARCHAR(100) NOT NULL,
            material_id VARCHAR(100) NULL,
            title VARCHAR(255) NOT NULL,
            source_file VARCHAR(255) NOT NULL,
            generated_at DATETIME NOT NULL,
            subject VARCHAR(255) NOT NULL DEFAULT 'General',
            main_topics LONGTEXT NOT NULL,
            key_concepts LONGTEXT NOT NULL,
            important_terms LONGTEXT NOT NULL,
            summary LONGTEXT NOT NULL,
            is_read TINYINT(1) NOT NULL DEFAULT 0,
            is_bookmarked TINYINT(1) NOT NULL DEFAULT 0,
            CONSTRAINT fk_reviewer_user FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE,
            CONSTRAINT fk_reviewer_material FOREIGN KEY (material_id) REFERENCES materials(id) ON DELETE SET NULL,
            INDEX idx_reviewer_user (user_id),
            INDEX idx_reviewer_material (material_id)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
    """)

    try:
        cur.execute("ALTER TABLE materials ADD COLUMN week VARCHAR(50) NOT NULL DEFAULT 'Week 1'")
    except Error:
        pass

    try:
        cur.execute("ALTER TABLE quizzes ADD COLUMN is_saved TINYINT(1) NOT NULL DEFAULT 0")
    except Error:
        pass

    try:
        cur.execute("ALTER TABLE materials ADD COLUMN folder_id VARCHAR(100) NULL")
    except Error:
        pass

    for statement in (
        "ALTER TABLE quizzes ADD COLUMN difficulty VARCHAR(20) NOT NULL DEFAULT 'normal'",
        "ALTER TABLE quiz_attempts ADD COLUMN difficulty VARCHAR(20) NOT NULL DEFAULT 'normal'",
    ):
        try:
            cur.execute(statement)
        except Error:
            pass

    # Add the expanded reviewer field to existing installations without
    # deleting or recreating the student's reviewer data.
    try:
        cur.execute("ALTER TABLE reviewers ADD COLUMN detailed_content LONGTEXT NULL")
    except Error:
        pass

    cur.execute("""
        CREATE TABLE IF NOT EXISTS quizzes (
            id VARCHAR(100) PRIMARY KEY,
            user_id VARCHAR(100) NOT NULL,
            material_id VARCHAR(100) NULL,
            reviewer_id VARCHAR(100) NULL,
            title VARCHAR(255) NOT NULL,
            source_type VARCHAR(30) NOT NULL DEFAULT 'material',
            source_title VARCHAR(255) NOT NULL,
            question_count INT NOT NULL DEFAULT 0,
            difficulty VARCHAR(20) NOT NULL DEFAULT 'normal',
            is_saved TINYINT(1) NOT NULL DEFAULT 0,
            created_at DATETIME NOT NULL,
            CONSTRAINT fk_quiz_user FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE,
            CONSTRAINT fk_quiz_material FOREIGN KEY (material_id) REFERENCES materials(id) ON DELETE SET NULL,
            CONSTRAINT fk_quiz_reviewer FOREIGN KEY (reviewer_id) REFERENCES reviewers(id) ON DELETE SET NULL,
            INDEX idx_quiz_user (user_id),
            INDEX idx_quiz_source (material_id, reviewer_id)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS quiz_questions (
            id VARCHAR(100) PRIMARY KEY,
            quiz_id VARCHAR(100) NOT NULL,
            question_order INT NOT NULL,
            question_text LONGTEXT NOT NULL,
            question_type VARCHAR(30) NOT NULL,
            choices_json LONGTEXT NULL,
            correct_answer LONGTEXT NOT NULL,
            explanation LONGTEXT NULL,
            CONSTRAINT fk_question_quiz FOREIGN KEY (quiz_id) REFERENCES quizzes(id) ON DELETE CASCADE,
            INDEX idx_question_quiz (quiz_id, question_order)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS quiz_attempts (
            id VARCHAR(100) PRIMARY KEY,
            user_id VARCHAR(100) NOT NULL,
            quiz_id VARCHAR(100) NULL,
            title VARCHAR(255) NOT NULL,
            source_type VARCHAR(30) NOT NULL DEFAULT 'material',
            source_title VARCHAR(255) NOT NULL,
            total_questions INT NOT NULL DEFAULT 0,
            difficulty VARCHAR(20) NOT NULL DEFAULT 'normal',
            correct_answers INT NOT NULL DEFAULT 0,
            score DECIMAL(6,2) NOT NULL DEFAULT 0,
            time_spent INT NOT NULL DEFAULT 0,
            completed_at DATETIME NOT NULL,
            CONSTRAINT fk_attempt_user FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE,
            CONSTRAINT fk_attempt_quiz FOREIGN KEY (quiz_id) REFERENCES quizzes(id) ON DELETE SET NULL,
            INDEX idx_attempt_user (user_id, completed_at),
            INDEX idx_attempt_quiz (quiz_id)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS quiz_answers (
            id VARCHAR(100) PRIMARY KEY,
            attempt_id VARCHAR(100) NOT NULL,
            question_id VARCHAR(100) NULL,
            question_text LONGTEXT NOT NULL,
            selected_answer LONGTEXT NULL,
            correct_answer LONGTEXT NOT NULL,
            is_correct TINYINT(1) NOT NULL DEFAULT 0,
            CONSTRAINT fk_answer_attempt FOREIGN KEY (attempt_id) REFERENCES quiz_attempts(id) ON DELETE CASCADE,
            CONSTRAINT fk_answer_question FOREIGN KEY (question_id) REFERENCES quiz_questions(id) ON DELETE SET NULL,
            INDEX idx_answer_attempt (attempt_id)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS notifications (
            id VARCHAR(100) PRIMARY KEY,
            user_id VARCHAR(100) NOT NULL,
            message LONGTEXT NOT NULL,
            type VARCHAR(30) NOT NULL DEFAULT 'info',
            read_flag TINYINT(1) NOT NULL DEFAULT 0,
            created_at DATETIME NOT NULL,
            CONSTRAINT fk_notification_user FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE,
            INDEX idx_notification_user (user_id, created_at)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS user_settings (
            user_id VARCHAR(100) PRIMARY KEY,
            dark_mode TINYINT(1) NOT NULL DEFAULT 0,
            offline_mode TINYINT(1) NOT NULL DEFAULT 0,
            notifications_enabled TINYINT(1) NOT NULL DEFAULT 1,
            auto_save TINYINT(1) NOT NULL DEFAULT 1,
            offline_access TINYINT(1) NOT NULL DEFAULT 1,
            updated_at DATETIME NOT NULL,
            CONSTRAINT fk_settings_user FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
    """)

    cur.execute("""
        CREATE TABLE IF NOT EXISTS ai_generations (
            id VARCHAR(100) PRIMARY KEY,
            user_id VARCHAR(100) NOT NULL,
            material_id VARCHAR(100) NULL,
            reviewer_id VARCHAR(100) NULL,
            generation_type VARCHAR(30) NOT NULL,
            model VARCHAR(100) NULL,
            status VARCHAR(30) NOT NULL,
            created_at DATETIME NOT NULL,
            completed_at DATETIME NULL,
            error_message LONGTEXT NULL,
            CONSTRAINT fk_ai_user FOREIGN KEY (user_id) REFERENCES users(id) ON DELETE CASCADE,
            CONSTRAINT fk_ai_material FOREIGN KEY (material_id) REFERENCES materials(id) ON DELETE SET NULL,
            CONSTRAINT fk_ai_reviewer FOREIGN KEY (reviewer_id) REFERENCES reviewers(id) ON DELETE SET NULL,
            INDEX idx_ai_user (user_id, created_at),
            INDEX idx_ai_material (material_id)
        ) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4
    """)

    conn.commit()
    cur.close()
    conn.close()


def _profile_from_row(row):
    # SELECT order: id, student_id, email, name, first_name, middle_name,
    # last_name, course, is_setup_complete, password_hash, is_guest.
    first = row[4] or ""
    middle = row[5] or ""
    last = row[6] or ""
    full = " ".join(x for x in (first, middle, last) if x).strip() or (row[3] or "")
    return {
        "name": full,
        "firstName": first,
        "middleName": middle,
        "lastName": last,
        "studentId": row[1] or "",
        "email": row[2] or "",
        "course": row[7] or "",
        "isSetupComplete": bool(row[8]),
    }


def _material_from_row(row):
    # SELECT order: id,user_id,title,original_filename,subject,topic,week,file_type,stored_filename,mime_type,file_path,size,uploaded_at,can_generate_ai
    return {
        "id": row[0],
        "filename": row[3],
        "type": row[7],
        "subject": row[4],
        "topic": row[5],
        "week": row[6] or "Week 1",
        "uploadedAt": row[12].strftime("%b %-d, %Y") if os.name != "nt" else row[12].strftime("%b %#d, %Y"),
        "size": f"{int(row[11]) / 1024 / 1024:.2f} MB",
        "canGenerateAI": bool(row[13]),
        "storedFilename": row[8],
        "mimeType": row[9] or "",
        "url": f"http://localhost:3001/uploads/{row[8]}",
        "filePath": row[10],
        "sizeBytes": int(row[11]),
        "folderId": row[14] if len(row) > 14 else None,
    }


def _reviewer_from_row(row):
    return {
        "id": row[0],
        "title": row[2],
        "sourceFile": row[3],
        "sourceMaterialId": row[1],
        "generatedAt": row[4].strftime("%b %-d, %Y") if os.name != "nt" else row[4].strftime("%b %#d, %Y"),
        "subject": row[5],
        "mainTopics": jload(row[6], []),
        "keyConcepts": jload(row[7], []),
        "importantTerms": jload(row[8], []),
        "summary": row[9],
        "detailedContent": jload(row[12], {}),
        "isRead": bool(row[10]),
        "isBookmarked": bool(row[11]),
    }


def _question_from_row(row):
    return {
        "id": row[0],
        "text": row[2],
        "type": row[3],
        "choices": jload(row[4], None),
        "correctAnswer": row[5],
        **({"explanation": row[6]} if row[6] else {}),
    }


def _settings_from_row(row):
    if not row:
        return DEFAULT_SETTINGS.copy()
    return {
        "darkMode": bool(row[1]),
        "offlineMode": bool(row[2]),
        "notificationsEnabled": bool(row[3]),
        "autoSave": bool(row[4]),
        "offlineAccess": bool(row[5]),
    }


def get_user(user_id):
    if not user_id:
        return None
    conn = mysql_connection()
    cur = conn.cursor()
    cur.execute("""
        SELECT id, student_id, email, name, first_name, middle_name, last_name, course, is_setup_complete,
               password_hash, is_guest
        FROM users WHERE id=%s
    """, (user_id,))
    row = cur.fetchone()
    if not row:
        cur.close(); conn.close()
        return None

    cur.execute("SELECT id, name, created_at FROM material_folders WHERE user_id=%s ORDER BY name", (user_id,))
    folders = [{"id": r[0], "name": r[1], "createdAt": r[2].isoformat()} for r in cur.fetchall()]

    cur.execute("""
        SELECT id, user_id, title, original_filename, subject, topic, week, file_type,
               stored_filename, mime_type, file_path, file_size_bytes,
               uploaded_at, can_generate_ai, folder_id
        FROM materials WHERE user_id=%s ORDER BY uploaded_at DESC
    """, (user_id,))
    materials = [_material_from_row(x) for x in cur.fetchall()]

    cur.execute("""
        SELECT id, material_id, title, source_file, generated_at, subject,
               main_topics, key_concepts, important_terms, summary,
               is_read, is_bookmarked, detailed_content
        FROM reviewers WHERE user_id=%s ORDER BY generated_at DESC
    """, (user_id,))
    reviewers = [_reviewer_from_row(x) for x in cur.fetchall()]

    cur.execute("""
        SELECT id, user_id, quiz_id, title, source_type, source_title,
               total_questions, difficulty, correct_answers, score, time_spent, completed_at
        FROM quiz_attempts WHERE user_id=%s ORDER BY completed_at DESC
    """, (user_id,))
    attempts = cur.fetchall()
    quiz_history = []
    for a in attempts:
        cur.execute("""
            SELECT id, quiz_id, question_order, question_text,
                   question_type, choices_json, correct_answer, explanation
            FROM quiz_questions
            WHERE quiz_id=%s ORDER BY question_order
        """, (a[2],)) if a[2] else None
        questions = [_question_from_row(x) for x in (cur.fetchall() if a[2] else [])]
        cur.execute("""
            SELECT question_id, selected_answer
            FROM quiz_answers WHERE attempt_id=%s
        """, (a[0],))
        answers = {str(x[0]): x[1] for x in cur.fetchall() if x[0]}
        cur.execute("SELECT COALESCE(is_saved,0) FROM quizzes WHERE id=%s", (a[2],)) if a[2] else None
        saved_row = cur.fetchone() if a[2] else None
        quiz_history.append({
            "id": a[0], "quizId": a[2], "isSaved": bool(saved_row[0]) if saved_row else False, "title": a[3], "sourceType": a[4],
            "sourceTitle": a[5], "totalQuestions": a[6], "difficulty": a[7] or "normal",
            "correctAnswers": a[8], "score": float(a[9]),
            "completedAt": a[11].strftime("%b %-d, %Y") if os.name != "nt" else a[11].strftime("%b %#d, %Y"),
            "timeSpent": int(a[10]),
            "questions": questions, "answers": answers,
        })

    cur.execute("""
        SELECT id, message, type, read_flag, created_at
        FROM notifications WHERE user_id=%s ORDER BY created_at DESC
    """, (user_id,))
    notifications = []
    for n in cur.fetchall():
        notifications.append({
            "id": n[0], "message": n[1], "type": n[2],
            "time": "Just now" if (utc_now() - n[4]).total_seconds() < 60 else n[4].strftime("%b %d"),
            "read": bool(n[3]),
        })

    cur.execute("""
        SELECT user_id, dark_mode, offline_mode, notifications_enabled,
               auto_save, offline_access, updated_at
        FROM user_settings WHERE user_id=%s
    """, (user_id,))
    settings = _settings_from_row(cur.fetchone())

    user = {
        "profile": _profile_from_row(row),
        "materials": materials,
        "folders": folders,
        "reviewers": reviewers,
        "quizHistory": quiz_history,
        "notifications": notifications,
        "settings": settings,
    }
    if row[9]:
        user["auth"] = {"passwordHash": row[9]}
    if row[10]:
        user["guest"] = True

    cur.close(); conn.close()
    return user


def ensure_user(user_id):
    if not user_id:
        return None
    user = get_user(user_id)
    if user:
        return user
    user = {
        "profile": {"name": "", "firstName": "", "middleName": "", "lastName": "", "studentId": "", "email": "", "course": "", "isSetupComplete": False},
        "materials": [], "reviewers": [], "quizHistory": [], "notifications": [],
        "settings": DEFAULT_SETTINGS.copy(),
    }
    conn = mysql_connection()
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO users (id, student_id, email, name, middle_name, last_name, course, is_setup_complete,
                           password_hash, is_guest, materials, reviewers, quiz_history,
                           notifications, settings)
        VALUES (%s,%s,%s,%s,%s,%s,%s,0,NULL,0,'[]','[]','[]','[]',%s)
    """, (user_id, "", "", "", "", "", "", jdump(DEFAULT_SETTINGS)))
    cur.execute("""
        INSERT IGNORE INTO user_settings
        (user_id, dark_mode, offline_mode, notifications_enabled, auto_save, offline_access, updated_at)
        VALUES (%s,0,0,1,1,1,%s)
    """, (user_id, utc_now()))
    conn.commit(); cur.close(); conn.close()
    return user


def save_profile_auth(user_id, user, password_value=None):
    profile = user.get("profile", {})
    auth = user.get("auth", {})
    first = (profile.get("firstName") or "").strip()
    middle = (profile.get("middleName") or "").strip()
    last = (profile.get("lastName") or "").strip()
    if not (first or middle or last):
        parts = (profile.get("name") or "").strip().split()
        first = parts[0] if parts else ""
        last = parts[-1] if len(parts) > 1 else ""
        middle = " ".join(parts[1:-1]) if len(parts) > 2 else ""
    full = " ".join(x for x in (first, middle, last) if x).strip()
    conn = mysql_connection(); cur = conn.cursor()
    cur.execute("""
        UPDATE users
        SET student_id=%s, email=%s, name=%s, first_name=%s, middle_name=%s,
            last_name=%s, course=%s, is_setup_complete=%s, password_hash=%s, is_guest=%s
        WHERE id=%s
    """, (
        profile.get("studentId", ""), profile.get("email", ""), full, first, middle, last,
        profile.get("course", ""), int(bool(profile.get("isSetupComplete"))),
        password_value if password_value is not None else auth.get("passwordHash"),
        int(bool(user.get("guest", False))), user_id
    ))
    conn.commit(); cur.close(); conn.close()


def save_settings(user_id, settings):
    s = {**DEFAULT_SETTINGS, **settings}
    conn = mysql_connection()
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO user_settings
        (user_id, dark_mode, offline_mode, notifications_enabled, auto_save, offline_access, updated_at)
        VALUES (%s,%s,%s,%s,%s,%s,%s)
        ON DUPLICATE KEY UPDATE
            dark_mode=VALUES(dark_mode),
            offline_mode=VALUES(offline_mode),
            notifications_enabled=VALUES(notifications_enabled),
            auto_save=VALUES(auto_save),
            offline_access=VALUES(offline_access),
            updated_at=VALUES(updated_at)
    """, (user_id, int(s["darkMode"]), int(s["offlineMode"]), int(s["notificationsEnabled"]),
          int(s["autoSave"]), int(s["offlineAccess"]), utc_now()))
    conn.commit(); cur.close(); conn.close()


def add_notification(user_id, message, ntype="success"):
    conn = mysql_connection()
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO notifications (id,user_id,message,type,read_flag,created_at)
        VALUES (%s,%s,%s,%s,0,%s)
    """, (uid("n-"), user_id, message, ntype, utc_now()))
    conn.commit(); cur.close(); conn.close()


def current_user_id():
    return request.headers.get("x-user-id", "")


def require_user():
    # Never create an account as a side effect of a normal API request.
    # Only explicit registration/guest flows may create users.
    user_id = current_user_id()
    return get_user(user_id) if user_id else None


def migrate_legacy_rows():
    """Move the original JSON-in-users data into normalized tables once per user."""
    conn = mysql_connection()
    cur = conn.cursor()
    cur.execute("""
        SELECT id, materials, reviewers, quiz_history, notifications, settings
        FROM users
    """)
    rows = cur.fetchall()
    for user_id, materials_json, reviewers_json, quiz_json, notifications_json, settings_json in rows:
        cur.execute("SELECT COUNT(*) FROM materials WHERE user_id=%s", (user_id,))
        materials_empty = cur.fetchone()[0] == 0
        if materials_empty:
            for m in jload(materials_json, []):
                stored = m.get("storedFilename", "")
                size_bytes = 0
                if stored and (UPLOAD_DIR / stored).exists():
                    size_bytes = (UPLOAD_DIR / stored).stat().st_size
                uploaded = utc_now()
                cur.execute("""
                    INSERT IGNORE INTO materials
                    (id,user_id,title,original_filename,subject,topic,file_type,mime_type,
                     stored_filename,file_path,file_size_bytes,uploaded_at,updated_at,can_generate_ai)
                    VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                """, (
                    m.get("id") or uid("m-"), user_id,
                    Path(m.get("filename") or m.get("title") or "Material").stem,
                    m.get("filename") or "Material",
                    m.get("subject") or "General", m.get("topic") or "General Topic",
                    m.get("type") or "IMAGE", m.get("mimeType"), stored,
                    str(UPLOAD_DIR / stored), size_bytes, uploaded, uploaded,
                    int(bool(m.get("canGenerateAI")))
                ))

        cur.execute("SELECT COUNT(*) FROM reviewers WHERE user_id=%s", (user_id,))
        if cur.fetchone()[0] == 0:
            for r in jload(reviewers_json, []):
                cur.execute("""
                    INSERT IGNORE INTO reviewers
                    (id,user_id,material_id,title,source_file,generated_at,subject,
                     main_topics,key_concepts,important_terms,summary,is_read,is_bookmarked)
                    VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
                """, (
                    r.get("id") or uid("r-"), user_id, r.get("sourceMaterialId"),
                    r.get("title") or "Reviewer", r.get("sourceFile") or "",
                    utc_now(), r.get("subject") or "General",
                    jdump(r.get("mainTopics", [])), jdump(r.get("keyConcepts", [])),
                    jdump(r.get("importantTerms", [])), r.get("summary") or "",
                    int(bool(r.get("isRead"))), int(bool(r.get("isBookmarked")))
                ))

        cur.execute("SELECT COUNT(*) FROM notifications WHERE user_id=%s", (user_id,))
        if cur.fetchone()[0] == 0:
            for n in jload(notifications_json, []):
                cur.execute("""
                    INSERT IGNORE INTO notifications
                    (id,user_id,message,type,read_flag,created_at)
                    VALUES (%s,%s,%s,%s,%s,%s)
                """, (
                    n.get("id") or uid("n-"), user_id, n.get("message") or "",
                    n.get("type") or "info", int(bool(n.get("read"))), utc_now()
                ))

        cur.execute("SELECT COUNT(*) FROM user_settings WHERE user_id=%s", (user_id,))
        if cur.fetchone()[0] == 0:
            s = {**DEFAULT_SETTINGS, **jload(settings_json, {})}
            cur.execute("""
                INSERT INTO user_settings
                (user_id,dark_mode,offline_mode,notifications_enabled,auto_save,offline_access,updated_at)
                VALUES (%s,%s,%s,%s,%s,%s,%s)
            """, (user_id, int(s["darkMode"]), int(s["offlineMode"]), int(s["notificationsEnabled"]),
                  int(s["autoSave"]), int(s["offlineAccess"]), utc_now()))

    conn.commit(); cur.close(); conn.close()
    print("[DATABASE] Legacy JSON data migration checked.")


def migrate_json_once():
    if not DB_FILE.exists():
        return
    try:
        old = json.loads(DB_FILE.read_text(encoding="utf-8"))
        old_users = old.get("users", {})
    except Exception:
        return
    if not old_users:
        return
    conn = mysql_connection()
    cur = conn.cursor()
    cur.execute("SELECT COUNT(*) FROM users")
    count = cur.fetchone()[0]
    cur.close(); conn.close()
    if count:
        return
    for user_id, user in old_users.items():
        profile = user.get("profile", {})
        auth = user.get("auth", {})
        conn = mysql_connection()
        cur = conn.cursor()
        cur.execute("""
            INSERT IGNORE INTO users
            (id,student_id,email,name,course,is_setup_complete,password_hash,is_guest,
             materials,reviewers,quiz_history,notifications,settings)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
        """, (
            user_id, profile.get("studentId",""), profile.get("email",""),
            profile.get("name",""), profile.get("course",""),
            int(bool(profile.get("isSetupComplete"))), auth.get("passwordHash"),
            int(bool(user.get("guest"))), jdump(user.get("materials",[])),
            jdump(user.get("reviewers",[])), jdump(user.get("quizHistory",[])),
            jdump(user.get("notifications",[])), jdump(user.get("settings",DEFAULT_SETTINGS))
        ))
        conn.commit(); cur.close(); conn.close()
    print(f"[DATABASE] Migrated {len(old_users)} legacy user(s) from db.json.")


@app.get("/api/health")
def health():
    try:
        conn = mysql_connection()
        cur = conn.cursor()
        cur.execute("SELECT 1")
        cur.fetchone()
        cur.close(); conn.close()
        return jsonify(ok=True, service="ISLearning Python backend", database="MySQL/MariaDB")
    except Error as e:
        return jsonify(ok=False, service="ISLearning Python backend",
                        database="MySQL/MariaDB", error=str(e)), 500


@app.post("/api/auth/register")
def register():
    data = request.get_json(silent=True) or {}
    student_id = (data.get("studentId") or "").strip()
    email = (data.get("email") or "").strip().lower()
    password = data.get("password") or ""
    if not student_id or not email or not password:
        return jsonify(message="Student ID, email, and password are required."), 400
    if len(password) < 6:
        return jsonify(message="Password must be at least 6 characters."), 400
    first_name = (data.get("firstName") or "").strip()
    middle_name = (data.get("middleName") or "").strip()
    last_name = (data.get("lastName") or "").strip()
    course = (data.get("course") or "").strip()
    setup_complete = bool(first_name and last_name and course)
    full_name = " ".join(part for part in (first_name, middle_name, last_name) if part)

    conn = mysql_connection()
    cur = conn.cursor()
    cur.execute("SELECT id FROM users WHERE student_id=%s OR LOWER(email)=LOWER(%s) LIMIT 1", (student_id, email))
    if cur.fetchone():
        cur.close(); conn.close()
        return jsonify(message="An account with that Student ID or email already exists."), 409

    user_id = uid("user-")
    cur.execute("""
        INSERT INTO users (id, student_id, email, name, first_name, middle_name, last_name, course, is_setup_complete, password_hash, is_guest,
                           materials, reviewers, quiz_history, notifications, settings)
        VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,0,'[]','[]','[]','[]',%s)
    """, (user_id, student_id, email, full_name, first_name, middle_name, last_name, course, int(setup_complete), password_hash(password), jdump(DEFAULT_SETTINGS)))
    cur.execute("""
        INSERT INTO user_settings
        (user_id, dark_mode, offline_mode, notifications_enabled, auto_save, offline_access, updated_at)
        VALUES (%s,0,0,1,1,1,%s)
    """, (user_id, utc_now()))
    conn.commit(); cur.close(); conn.close()
    user = get_user(user_id)
    return jsonify(userId=user_id, profile=user["profile"], message="Account created successfully."), 201


@app.post("/api/auth/login")
def login():
    data = request.get_json(silent=True) or {}
    identifier = (data.get("identifier") or "").strip()
    password = data.get("password") or ""
    if not identifier or not password:
        return jsonify(message="Student ID/email and password are required."), 400

    conn = mysql_connection()
    cur = conn.cursor()
    # Do not use LIMIT 1 here. Older ISLearning databases may contain duplicate
    # email/student identifiers. The previous query could select the wrong
    # account and report "Invalid credentials" even when the password was correct.
    # Check every matching account and use the one whose password verifies.
    identifier_email = identifier.lower()
    cur.execute("""
        SELECT id, student_id, email, password_hash
        FROM users
        WHERE student_id=%s OR LOWER(email)=LOWER(%s)
        ORDER BY CASE WHEN student_id=%s THEN 0 ELSE 1 END, created_at ASC
    """, (identifier, identifier_email, identifier))
    rows = cur.fetchall()
    if not rows:
        cur.close(); conn.close()
        return jsonify(message="Invalid credentials."), 401

    found = None
    expected = None
    for candidate in rows:
        if verify_password(password, candidate[3]):
            found = candidate[0]
            expected = candidate[3]
            break

    if not found:
        cur.close(); conn.close()
        return jsonify(message="Invalid credentials."), 401

    # Upgrade old SHA-256 passwords after successful login.
    if expected and not expected.startswith("pbkdf2_sha256$"):
        cur.execute("UPDATE users SET password_hash=%s WHERE id=%s", (password_hash(password), found))

    # Repair legacy accounts whose profile fields are present but the completion
    # flag was never persisted. Such returning users should not be sent through
    # onboarding again. If only the legacy full-name field exists, split it once.
    cur.execute("SELECT first_name, middle_name, last_name, name, course, student_id FROM users WHERE id=%s", (found,))
    profile_row = cur.fetchone()
    if profile_row:
        first_name, middle_name, last_name, legacy_name, course_value, sid_value = profile_row
        if not (first_name or '').strip() and (legacy_name or '').strip():
            parts = (legacy_name or '').strip().split()
            if len(parts) >= 2:
                first_name, last_name = parts[0], parts[-1]
                if len(parts) > 2 and not (middle_name or '').strip(): middle_name = ' '.join(parts[1:-1])
                cur.execute("UPDATE users SET first_name=%s, middle_name=%s, last_name=%s WHERE id=%s", (first_name, middle_name or '', last_name, found))
        if (first_name or '').strip() and (last_name or '').strip() and (course_value or '').strip() and (sid_value or '').strip():
            cur.execute("UPDATE users SET is_setup_complete=1 WHERE id=%s AND is_setup_complete=0", (found,))

    # Record the successful login and persist any legacy-profile repair.
    cur.execute("UPDATE users SET updated_at=CURRENT_TIMESTAMP WHERE id=%s", (found,))
    conn.commit(); cur.close(); conn.close()

    user = get_user(found)
    return jsonify(userId=found, profile=user["profile"])


@app.post("/api/auth/guest")
def guest():
    user_id = uid("guest-")
    user = ensure_user(user_id)
    user["profile"] = {
        "name": "Student", "firstName": "Student", "middleName": "", "lastName": "",
        "studentId": "", "email": "", "course": "", "isSetupComplete": True
    }
    user["guest"] = True
    save_profile_auth(user_id, user)
    return jsonify(userId=user_id, profile=user["profile"])


@app.get("/api/state")
def state():
    u = require_user()
    if not u:
        return jsonify(message="Missing x-user-id header."), 401
    return jsonify(u)


@app.put("/api/profile")
def profile():
    u = require_user()
    if not u:
        return jsonify(message="Missing x-user-id header."), 401
    u["profile"].update(request.get_json(silent=True) or {})
    save_profile_auth(current_user_id(), u)
    return jsonify(u["profile"])


@app.put("/api/settings")
def settings():
    u = require_user()
    if not u:
        return jsonify(message="Missing x-user-id header."), 401
    u["settings"].update(request.get_json(silent=True) or {})
    save_settings(current_user_id(), u["settings"])
    return jsonify(u["settings"])



SUPPORTED_AI_EXTENSIONS = {
    ".pdf": "PDF", ".doc": "DOC", ".docx": "DOCX", ".ppt": "PPT", ".pptx": "PPTX",
    ".txt": "TXT", ".md": "TXT", ".csv": "CSV", ".html": "HTML", ".xml": "XML",
    ".json": "JSON", ".xlsx": "XLSX", ".xls": "XLS",
    ".png": "IMAGE", ".jpg": "IMAGE", ".jpeg": "IMAGE", ".webp": "IMAGE",
}


def gemini_client():
    if not GEMINI_API_KEY or GEMINI_API_KEY == "PASTE_YOUR_GEMINI_API_KEY_HERE":
        raise RuntimeError("GEMINI_API_KEY is not configured. Add your own Gemini API key to backend/.env; hybrid mode will then use the local fallback if Gemini is unavailable.")
    return genai.Client(api_key=GEMINI_API_KEY)


def _is_retryable_gemini_error(exc):
    """Return True for errors where another current Gemini model may succeed."""
    message = str(exc).upper()
    markers = (
        "503", "UNAVAILABLE", "429", "RESOURCE_EXHAUSTED",
        "TOO MANY REQUESTS", "HIGH DEMAND", "RATE LIMIT",
        "MODEL_NOT_FOUND", "NOT_FOUND", "NOT FOUND",
    )
    return any(marker in message for marker in markers)


def _gemini_models_to_try():
    models = []
    for model in [GEMINI_MODEL, *GEMINI_FALLBACK_MODELS]:
        if model and model not in models:
            models.append(model)
    return models


# Gemini Files are reusable for up to 48 hours. Cache successful uploads in
# this process for 20 hours so repeat reviewer/quiz generation doesn't upload
# the same source file again. Cache keys use file content, not only its name,
# so updated files never accidentally reuse an old upload.
_GEMINI_FILE_CACHE = {}
_GEMINI_FILE_CACHE_LOCK = threading.Lock()
_GEMINI_FILE_CACHE_TTL = 20 * 60 * 60
_GEMINI_FILE_CACHE_LIMIT = 64


def _gemini_uploaded_file(client, path, mime_type):
    path = Path(path)
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    cache_key = (digest.hexdigest(), mime_type or "")
    now = time.time()
    with _GEMINI_FILE_CACHE_LOCK:
        cached = _GEMINI_FILE_CACHE.get(cache_key)
        if cached and now - cached[0] < _GEMINI_FILE_CACHE_TTL:
            return cached[1]
        if cached:
            _GEMINI_FILE_CACHE.pop(cache_key, None)

    uploaded = client.files.upload(
        file=str(path),
        config=types.UploadFileConfig(mime_type=mime_type) if mime_type else None,
    )
    with _GEMINI_FILE_CACHE_LOCK:
        _GEMINI_FILE_CACHE[cache_key] = (now, uploaded)
        # Bound memory for long-running local/demo sessions.
        while len(_GEMINI_FILE_CACHE) > _GEMINI_FILE_CACHE_LIMIT:
            oldest_key = min(_GEMINI_FILE_CACHE, key=lambda key: _GEMINI_FILE_CACHE[key][0])
            _GEMINI_FILE_CACHE.pop(oldest_key, None)
    return uploaded


def _extract_local_source_text(path):
    """Extract text for the local model without changing the original uploaded file."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext in (".txt", ".md", ".csv", ".json", ".html", ".css", ".js", ".xml"):
        return path.read_text(encoding="utf-8", errors="replace")
    if ext == ".pdf":
        from pypdf import PdfReader
        return "\n\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if ext == ".docx":
        from docx import Document
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            parts.append("\n".join(" | ".join(c.text for c in row.cells) for row in table.rows))
        return "\n\n".join(parts)
    if ext == ".pptx":
        from pptx import Presentation
        return "\n\n".join(
            "\n".join(sh.text for sh in slide.shapes if getattr(sh, "has_text_frame", False) and sh.text.strip())
            for slide in Presentation(str(path)).slides
        )
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(str(path), read_only=True, data_only=True)
        try:
            return "\n\n".join(
                "Sheet: " + ws.title + "\n" + "\n".join(" | ".join(str(v or "") for v in row) for row in ws.iter_rows(values_only=True))
                for ws in wb.worksheets
            )
        finally:
            wb.close()
    raise RuntimeError(
        f"Local AI cannot extract text from {ext or 'this file type'}. "
        "Use a text-based PDF, DOCX, PPTX, XLSX, or TXT file; scanned PDFs and images need OCR."
    )


def _ollama_json(path, prompt, schema):
    """Run schema-constrained JSON generation against a local Ollama server."""
    source_text = _extract_local_source_text(path).strip()
    if not source_text:
        raise RuntimeError("Local AI could not extract readable text from this file. Scanned PDFs/images need OCR.")
    # Keep source material complete where practical; fail clearly rather than silently truncating it.
    max_chars = int(os.getenv("OLLAMA_MAX_SOURCE_CHARS", "50000"))
    if len(source_text) > max_chars:
        raise RuntimeError(
            f"The source has {len(source_text):,} characters, exceeding the local fallback limit "
            f"({max_chars:,}). Increase OLLAMA_MAX_SOURCE_CHARS or use a model with a larger context window."
        )
    payload = {
        "model": OLLAMA_MODEL,
        "prompt": prompt + "\n\nSOURCE MATERIAL (authoritative; use only this content):\n" + source_text +
                  "\n\nReturn only a valid JSON object matching this schema:\n" + json.dumps(schema, ensure_ascii=False),
        "format": schema,
        "stream": False,
        "options": {"temperature": 0.2, "num_ctx": int(os.getenv("OLLAMA_NUM_CTX", "16384"))},
    }
    req = urllib.request.Request(
        OLLAMA_BASE_URL + "/api/generate",
        data=json.dumps(payload).encode("utf-8"),
        headers={"Content-Type": "application/json"},
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=OLLAMA_TIMEOUT_SECONDS) as response:
            body = json.loads(response.read().decode("utf-8"))
        raw = (body.get("response") or "").strip()
        if not raw:
            raise RuntimeError("Ollama returned an empty response.")
        result = json.loads(raw)
        if not isinstance(result, dict):
            raise RuntimeError("Ollama response was not a JSON object.")
        app.logger.info("Local Ollama generation completed using %s", OLLAMA_MODEL)
        return result
    except urllib.error.URLError as exc:
        reason = getattr(exc, "reason", exc)
        raise RuntimeError(
            f"Local AI fallback is unavailable at {OLLAMA_BASE_URL}. Install/start Ollama, "
            f"download model '{OLLAMA_MODEL}', then retry. Details: {reason}"
        ) from exc


def _gemini_json(path, mime_type, prompt, schema, max_output_tokens=60000):
    """Try Gemini first; when it is unavailable/quota-limited, use local Ollama."""
    last_error = None
    if AI_PROVIDER_MODE != "local":
        try:
            client = gemini_client()
            upload_started = time.perf_counter()
            uploaded = _gemini_uploaded_file(client, path, mime_type)
            app.logger.info("Gemini source ready in %.2fs (%s)", time.perf_counter() - upload_started, Path(path).name)
            models = _gemini_models_to_try()
            for model_index, model in enumerate(models):
                # Temporary 503/high-demand responses are common; retry each model with
                # exponential backoff before moving to the next configured model.
                attempts = 3
                for attempt in range(attempts):
                    generation_started = time.perf_counter()
                    try:
                        response = client.models.generate_content(
                            model=model,
                            contents=[uploaded, prompt],
                            config=types.GenerateContentConfig(
                                response_mime_type="application/json",
                                response_schema=schema,
                                temperature=0.2,
                                max_output_tokens=max_output_tokens,
                            ),
                        )
                        raw = (response.text or "").strip()
                        if not raw:
                            raise RuntimeError(f"Gemini model {model} returned an empty response.")
                        app.logger.info("Gemini generation completed in %.2fs using %s", time.perf_counter() - generation_started, model)
                        return json.loads(raw)
                    except Exception as exc:
                        last_error = exc
                        msg = str(exc).lower()
                        schema_error = any(x in msg for x in ("response_schema", "schema", "structured output", "invalid argument"))
                        if schema_error:
                            try:
                                response = client.models.generate_content(
                                    model=model,
                                    contents=[uploaded, prompt + "\nReturn ONLY valid JSON. Do not wrap it in markdown fences."],
                                    config=types.GenerateContentConfig(
                                        response_mime_type="application/json", temperature=0.2,
                                        max_output_tokens=max_output_tokens,
                                    ),
                                )
                                raw = (response.text or "").strip()
                                if raw.startswith("```"):
                                    raw = raw.strip("`")
                                    if raw.startswith("json"):
                                        raw = raw[4:].lstrip()
                                parsed = json.loads(raw)
                                if isinstance(parsed, dict):
                                    return parsed
                            except Exception as fallback_exc:
                                last_error = fallback_exc
                        if not _is_retryable_gemini_error(exc):
                            raise RuntimeError(f"Gemini model {model} failed: {exc}") from exc
                        if attempt + 1 < attempts:
                            time.sleep(1.5 * (2 ** attempt))
                            continue
                        break
            tried = ", ".join(models)
            last_error = RuntimeError(f"Gemini unavailable across models ({tried}). Last error: {last_error}")
        except Exception as exc:
            last_error = exc

    if AI_PROVIDER_MODE in ("hybrid", "local"):
        app.logger.warning("Using local AI fallback after Gemini failure: %s", last_error or "local-only mode")
        try:
            result = _ollama_json(path, prompt, schema)
            result["_aiProvider"] = "ollama-local"
            return result
        except Exception as local_exc:
            raise RuntimeError(
                "AI generation is temporarily unavailable. Gemini could not complete the request after retrying "
                "available models, and the local Ollama fallback is not reachable or could not process this file. "
                "Check the Gemini API key/model availability or start Ollama and download the configured local model. "
                f"Gemini details: {last_error or 'skipped by configuration'}. Local fallback details: {local_exc}"
            ) from local_exc

    raise RuntimeError(f"Gemini generation failed: {last_error or 'Gemini is disabled.'}") from last_error

def material_record(material_id, user_id):
    conn = mysql_connection(); cur = conn.cursor()
    cur.execute("""
        SELECT id,title,original_filename,subject,topic,week,file_type,mime_type,stored_filename,file_path
        FROM materials WHERE id=%s AND user_id=%s
    """, (material_id, user_id))
    row = cur.fetchone(); cur.close(); conn.close()
    if not row:
        return None
    return {
        "id": row[0], "title": row[1], "filename": row[2], "subject": row[3], "topic": row[4], "week": row[5],
        "fileType": row[6], "mimeType": row[7], "storedFilename": row[8], "filePath": row[9]
    }


def ai_source_for(source_type, source_id, user_id):
    if source_type == "reviewer":
        conn = mysql_connection(); cur = conn.cursor()
        cur.execute("SELECT material_id FROM reviewers WHERE id=%s AND user_id=%s", (source_id, user_id))
        row = cur.fetchone(); cur.close(); conn.close()
        if not row or not row[0]:
            raise ValueError("The selected reviewer is not linked to an uploaded material.")
        return material_record(row[0], user_id)
    return material_record(source_id, user_id)


REVIEWER_SCHEMA = {
    "type": "object",
    "properties": {
        "title": {"type": "string"},
        "mainTopics": {"type": "array", "items": {"type": "string"}},
        "keyConcepts": {"type": "array", "items": {"type": "string"}},
        "importantTerms": {"type": "array", "items": {
            "type": "object", "properties": {
                "term": {"type": "string"},
                "definition": {"type": "string"}
            },
            "required": ["term", "definition"]
        }},
        "summary": {"type": "string"},
        "detailedContent": {
            "type": "object",
            "properties": {
                "overview": {"type": "string"},
                "learningObjectives": {"type": "array", "items": {"type": "string"}},
                "sections": {"type": "array", "items": {
                    "type": "object", "properties": {
                        "heading": {"type": "string"},
                        "explanation": {"type": "string"},
                        "points": {"type": "array", "items": {"type": "string"}},
                        "subsections": {"type": "array", "items": {
                            "type": "object", "properties": {
                                "heading": {"type": "string"},
                                "explanation": {"type": "string"},
                                "points": {"type": "array", "items": {"type": "string"}},
                                "examples": {"type": "array", "items": {"type": "string"}}
                            },
                            "required": ["heading", "explanation", "points", "examples"]
                        }},
                        "examples": {"type": "array", "items": {"type": "string"}},
                        "importantNotes": {"type": "array", "items": {"type": "string"}}
                    },
                    "required": ["heading", "explanation", "points", "subsections", "examples", "importantNotes"]
                }},
                "comparisons": {"type": "array", "items": {
                    "type": "object", "properties": {
                        "title": {"type": "string"},
                        "columns": {"type": "array", "items": {"type": "string"}},
                        "rows": {"type": "array", "items": {"type": "array", "items": {"type": "string"}}}
                    },
                    "required": ["title", "columns", "rows"]
                }},
                "keyTakeaways": {"type": "array", "items": {"type": "string"}},
                "practiceQuestions": {"type": "array", "items": {
                    "type": "object", "properties": {
                        "type": {"type": "string"},
                        "question": {"type": "string"},
                        "answer": {"type": "string"}
                    },
                    "required": ["type", "question", "answer"]
                }}
            },
            "required": ["overview", "learningObjectives", "sections", "comparisons", "keyTakeaways", "practiceQuestions"]
        }
    },
    "required": ["title", "mainTopics", "keyConcepts", "importantTerms", "summary", "detailedContent"]
}

QUIZ_SCHEMA = {
    "type": "object",
    "properties": {
        "title": {"type": "string"},
        "questions": {"type": "array", "items": {
            "type": "object",
            "properties": {
                "text": {"type": "string"},
                "type": {"type": "string", "enum": ["multiple-choice", "true-false", "enumeration", "identification", "practical-coding", "code-analysis", "expected-output"]},
                "choices": {"type": "array", "items": {"type": "string"}},
                "correctAnswer": {"type": "string"},
                "explanation": {"type": "string"}
            },
            "required": ["text", "type", "correctAnswer", "explanation"]
        }}
    },
    "required": ["title", "questions"]
}


def generate_reviewer_from_material(material):
    path = Path(material["filePath"])
    if not path.exists():
        raise FileNotFoundError("The uploaded source file is missing from the server.")
    prompt = f"""
You are the full academic reviewer generator inside ISLearning.

READ AND USE THE ENTIRE UPLOADED LEARNING MATERIAL from beginning to end. The uploaded file is the authoritative
source. Do not treat the task as a keyword extractor. Do not summarize away sections simply because they are long.
Do not use outside knowledge and do not invent facts.

Create a COMPLETE STUDY REVIEWER that a student could use to study the source without losing the important content.
Preserve the source's organization, terminology, order, definitions, explanations, examples, procedures, lists,
comparisons, tables, learning objectives, takeaways, and examination/practice material when those are present.

STYLE TARGET (match the academic, Word-ready structure requested in the user's Research_Reviewer_Word_Ready example):
- Make it look and read like a polished, Word-ready academic reviewer, not a short AI summary or a collection of cards.
- Start with a clear title such as "Reviewer: [actual source topic]".
- Use organized Roman-numeral major sections (I, II, III...) when suitable, numbered subtopics, descriptive headings,
  concise explanatory paragraphs, and bullet lists for classifications, steps, characteristics, and examples.
- Include learning objectives when the source supports them; definitions; types/classifications; comparison tables;
  examples; important distinctions; key takeaways; important terms to remember; and possible examination questions
  with answers when these are present or directly supportable from the source.
- Preserve the source's progression and important named terms. Do not force irrelevant sections into the output if the
  uploaded material does not contain that kind of content.
- Detailed content must be visible in the generated reviewer, not hidden only in a summary field.

SOURCE ACCURACY IS MORE IMPORTANT THAN POLISH:
- Treat only the uploaded file(s) as authoritative. Do not use outside knowledge to fill gaps or silently correct the source.
- Do not fabricate facts, definitions, examples, citations, equations, names, dates, or answers.
- Keep source-specific terms and claims faithful. If something is unclear or incomplete in the source, do not guess.
- For practice questions, only include an answer when it is explicitly given or clearly supported elsewhere in the same source.
- Cover all meaningful sections of the source, including later pages/slides; do not over-focus on the beginning.

IMPORTANT COVERAGE RULES:
- Identify every meaningful lesson/topic/heading in the source and represent it in detailedContent.sections.
- Keep sections in the same logical order as the source.
- For each section, explain what the source says, not merely its title.
- Preserve important subtopics as subsections.
- Preserve lists of characteristics, steps, types, classifications, guidelines, advantages/disadvantages, or criteria.
- Preserve examples when the source gives examples.
- Convert important source comparisons/tables into comparisons when useful, without changing their meaning.
- Include important definitions and relationships in the detailed sections even if they also appear in importantTerms.
- If the source contains examination questions, practice questions, identification, true/false, or short-answer material,
  include those in practiceQuestions with the answer supported by the source.
- Do not invent missing answers. If the source gives questions without answers, answer only when the answer is directly
  supported elsewhere in the same uploaded material.
- Do not arbitrarily limit the number of sections, concepts, terms, or questions. Coverage is more important than brevity.
- The overview and summary may be concise, but detailedContent.sections must be comprehensive.
- Do not copy huge passages verbatim. Paraphrase faithfully while preserving all important meaning.

Subject: {material['subject']}
User-provided topic: {material['topic']}
Source filename: {material['filename']}

OUTPUT REQUIREMENTS:
- mainTopics: a list of the major topics actually present in the source.
- keyConcepts: the important study concepts drawn from the whole source.
- importantTerms: important terms from the source with accurate definitions.
- summary: a coherent summary of the whole source.
- detailedContent.overview: a short orientation to the entire material.
- detailedContent.learningObjectives: source-supported learning objectives, or [] if the source has none.
- detailedContent.sections: the COMPLETE ordered lesson-by-lesson study guide.
- detailedContent.comparisons: important source-based comparisons/tables; use [] if none exist.
- detailedContent.keyTakeaways: important takeaways from across the source.
- detailedContent.practiceQuestions: useful study questions from across the source, with source-supported answers.

Return JSON matching the required schema.
"""
    return _gemini_json(path, material["mimeType"], prompt, REVIEWER_SCHEMA)

def generate_quiz_from_material(material, count, question_type="mixed", difficulty="normal", selected_types=None):
    path = Path(material["filePath"])
    if not path.exists():
        raise FileNotFoundError("The uploaded source file is missing from the server.")
    count = max(1, min(int(count or 10), 200))
    selected_types = selected_types or ([question_type] if question_type != "mixed" else ["multiple-choice", "true-false"])
    type_text = ", ".join(selected_types)
    prompt = f"""
Create a source-grounded practice quiz from the ENTIRE uploaded learning material.
Read the complete source, not just the first page or obvious keywords. Questions must be supported by the source.
Treat the uploaded source material as authoritative and do not add facts from outside knowledge. Do not invent definitions,
examples, terminology, dates, procedures, or claims. Ensure questions cover the full source instead of repeatedly testing
only the first sections. Before returning the quiz, verify that each question is answerable from the source, each answer
is explicitly supported by the source, and each explanation accurately reflects that source. Remove any ambiguous or
unsupported item. Difficulty: {difficulty}. Generate exactly {count} questions when possible.
Allowed question types (use ONLY these selected types): {type_text}.
Question-type compliance is required. Do not silently turn enumeration, identification, practical-coding, code-analysis, or expected-output questions into multiple-choice questions. For each selected type, create questions in that exact format when the requested count permits; distribute the requested count across selected types. If only one type is selected, every question must use that type.
For enumeration, ask the learner to list the required items and put the complete expected list in correctAnswer, separated by semicolons. For identification, ask for a term/concept and provide the exact expected term. For practical-coding, require a code solution. For code-analysis, show a short source-grounded code snippet in the question and ask the learner to explain/analyze it. For expected-output, include the complete code snippet and ask for its output. For true-false, correctAnswer must be exactly True or False. For multiple-choice, choices must contain 3-4 distinct options and correctAnswer must exactly match one choice.

Difficulty rules:
- easy: direct definitions, recall, simple applications.
- normal: application, moderate reasoning, multi-step understanding.
- hard: deeper analysis, troubleshooting, complex application.

For multiple-choice, provide 3-4 choices and set correctAnswer to the exact correct choice.
For true-false, choices may be ["True","False"].
For enumeration/identification/coding/open-ended types, correctAnswer should contain the expected answer or rubric and the explanation should say that manual or AI-assisted checking may be needed.
Return only JSON matching the schema.
"""
    generated = _gemini_json(
        path, material.get("mimeType"), prompt, QUIZ_SCHEMA,
        max_output_tokens=min(20000, max(6000, count * 750 + 1500))
    )
    questions = generated.get("questions") or []
    normalized = []
    allowed = {"multiple-choice","true-false","enumeration","identification","practical-coding","code-analysis","expected-output"}
    for q in questions[:count]:
        qt = q.get("type", "multiple-choice")
        if qt not in allowed:
            qt = "multiple-choice"
        normalized.append({
            "id": uid("q-"), "text": str(q.get("text", "")).strip(), "type": qt,
            "choices": q.get("choices") or ([] if qt not in ("multiple-choice","true-false") else (["True","False"] if qt == "true-false" else [])),
            "correctAnswer": str(q.get("correctAnswer", "")).strip(),
            "explanation": str(q.get("explanation", "")).strip(),
        })
    if len(normalized) < min(count, 5):
        raise RuntimeError("Gemini returned too few usable quiz questions. Please try again.")
    return generated.get("title") or f"{Path(material['filename']).stem} Practice Quiz", normalized


@app.get("/api/folders")
def list_folders():
    if not require_user():
        return jsonify(message="Missing x-user-id header."), 401
    conn = mysql_connection(); cur = conn.cursor()
    cur.execute("SELECT id, name, created_at FROM material_folders WHERE user_id=%s ORDER BY name", (current_user_id(),))
    rows = cur.fetchall(); cur.close(); conn.close()
    return jsonify([{ "id": r[0], "name": r[1], "createdAt": r[2].isoformat() } for r in rows])


@app.post("/api/folders")
def create_folder():
    if not require_user():
        return jsonify(message="Missing x-user-id header."), 401
    name = (request.get_json(silent=True) or {}).get("name", "").strip()
    if not name or len(name) > 255:
        return jsonify(message="Enter a folder name up to 255 characters."), 400
    folder_id = uid("folder-")
    conn = mysql_connection(); cur = conn.cursor()
    try:
        cur.execute("INSERT INTO material_folders (id,user_id,name,created_at) VALUES (%s,%s,%s,%s)", (folder_id, current_user_id(), name, utc_now()))
        conn.commit()
    except Error as exc:
        conn.rollback(); cur.close(); conn.close()
        return jsonify(message="A folder with that name already exists."), 409
    cur.close(); conn.close()
    return jsonify(id=folder_id, name=name), 201


@app.patch("/api/folders/<folder_id>")
def rename_folder(folder_id):
    if not require_user():
        return jsonify(message="Missing x-user-id header."), 401
    name = (request.get_json(silent=True) or {}).get("name", "").strip()
    if not name or len(name) > 255:
        return jsonify(message="Enter a folder name up to 255 characters."), 400
    conn = mysql_connection(); cur = conn.cursor()
    try:
        cur.execute("UPDATE material_folders SET name=%s WHERE id=%s AND user_id=%s", (name, folder_id, current_user_id()))
        if not cur.rowcount:
            conn.rollback(); cur.close(); conn.close(); return jsonify(message="Folder not found."), 404
        conn.commit()
    except Error:
        conn.rollback(); cur.close(); conn.close(); return jsonify(message="A folder with that name already exists."), 409
    cur.close(); conn.close(); return jsonify(id=folder_id, name=name)


@app.delete("/api/folders/<folder_id>")
def delete_folder(folder_id):
    if not require_user():
        return jsonify(message="Missing x-user-id header."), 401
    conn = mysql_connection(); cur = conn.cursor()
    cur.execute("UPDATE materials SET folder_id=NULL WHERE folder_id=%s AND user_id=%s", (folder_id, current_user_id()))
    cur.execute("DELETE FROM material_folders WHERE id=%s AND user_id=%s", (folder_id, current_user_id()))
    if not cur.rowcount:
        conn.rollback(); cur.close(); conn.close(); return jsonify(message="Folder not found."), 404
    conn.commit(); cur.close(); conn.close(); return jsonify(ok=True)


@app.patch("/api/materials/<mid>/folder")
def move_material_to_folder(mid):
    if not require_user():
        return jsonify(message="Missing x-user-id header."), 401
    folder_id = (request.get_json(silent=True) or {}).get("folderId") or None
    conn = mysql_connection(); cur = conn.cursor()
    if folder_id:
        cur.execute("SELECT id FROM material_folders WHERE id=%s AND user_id=%s", (folder_id, current_user_id()))
        if not cur.fetchone():
            cur.close(); conn.close(); return jsonify(message="Folder not found."), 404
    cur.execute("UPDATE materials SET folder_id=%s, updated_at=%s WHERE id=%s AND user_id=%s", (folder_id, utc_now(), mid, current_user_id()))
    if not cur.rowcount:
        conn.rollback(); cur.close(); conn.close(); return jsonify(message="Material not found."), 404
    conn.commit(); cur.close(); conn.close()
    return jsonify(id=mid, folderId=folder_id)


@app.post("/api/materials/upload")
def upload_material():
    u = require_user()
    if not u:
        return jsonify(message="Missing x-user-id header."), 401
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify(message="Please select a file."), 400

    original = file.filename
    safe = secure_filename(original) or "upload"
    ext = Path(original).suffix.lower()
    typ = SUPPORTED_AI_EXTENSIONS.get(ext, "IMAGE")
    # Keep uploads organized by account and file category instead of placing every
    # uploaded file in the same directory. The database stores the relative path.
    user_folder = secure_filename(str(current_user_id())) or "user"
    category = {"PDF": "documents", "DOCX": "documents", "PPTX": "presentations", "TXT": "text", "XLSX": "spreadsheets"}.get(typ, "images")
    stored_name = f"{int(datetime.now().timestamp()*1000)}-{uuid.uuid4().hex[:10]}-{safe}"
    relative_stored = str(Path(user_folder) / category / stored_name).replace("\\", "/")
    target = UPLOAD_DIR / relative_stored
    target.parent.mkdir(parents=True, exist_ok=True)
    file.save(target)

    size = target.stat().st_size
    material_id = uid("m-")
    now = utc_now()

    conn = mysql_connection()
    cur = conn.cursor()
    cur.execute("""
        INSERT INTO materials
        (id,user_id,title,original_filename,subject,topic,week,file_type,mime_type,
         stored_filename,file_path,file_size_bytes,uploaded_at,updated_at,can_generate_ai)
        VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
    """, (
        material_id, current_user_id(), Path(original).stem, original,
        request.form.get("subject") or "General",
        request.form.get("topic") or "General Topic",
        request.form.get("week") or "Week 1", typ, file.mimetype,
        relative_stored, str(target), size, now, now, int(typ != "IMAGE")
    ))
    conn.commit(); cur.close(); conn.close()

    add_notification(current_user_id(), f"Material uploaded: {original}")
    material = get_user(current_user_id())["materials"][0]
    return jsonify(material), 201


@app.patch("/api/materials/<mid>/organize")
def organize_material(mid):
    if not require_user():
        return jsonify(message="Missing x-user-id header."), 401
    data = request.get_json(silent=True) or {}
    subject = (data.get("subject") or "").strip()
    week = (data.get("week") or "").strip()
    topic = (data.get("topic") or "").strip()
    if not subject or not week:
        return jsonify(message="Subject and week are required."), 400
    conn = mysql_connection(); cur = conn.cursor()
    cur.execute("""UPDATE materials SET subject=%s, week=%s, topic=COALESCE(NULLIF(%s,''),topic), updated_at=%s
                   WHERE id=%s AND user_id=%s""", (subject, week, topic, utc_now(), mid, current_user_id()))
    changed = cur.rowcount
    conn.commit(); cur.close(); conn.close()
    if not changed:
        return jsonify(message="Material not found."), 404
    return jsonify(next(x for x in get_user(current_user_id())["materials"] if x["id"] == mid))

@app.patch("/api/materials/<mid>")
def rename_material(mid):
    if not require_user():
        return jsonify(message="Missing x-user-id header."), 401
    data = request.get_json(silent=True) or {}
    new_name = (data.get("filename") or "").strip()
    if not new_name:
        return jsonify(message="Filename is required."), 400

    conn = mysql_connection()
    cur = conn.cursor()
    cur.execute("""
        UPDATE materials SET original_filename=%s, title=%s, updated_at=%s
        WHERE id=%s AND user_id=%s
    """, (new_name, Path(new_name).stem, utc_now(), mid, current_user_id()))
    changed = cur.rowcount
    conn.commit(); cur.close(); conn.close()
    if not changed:
        return jsonify(message="Material not found."), 404
    return jsonify(next(x for x in get_user(current_user_id())["materials"] if x["id"] == mid))


@app.delete("/api/materials/<mid>")
def delete_material(mid):
    if not require_user():
        return jsonify(message="Missing x-user-id header."), 401

    conn = mysql_connection()
    cur = conn.cursor()
    cur.execute("SELECT stored_filename FROM materials WHERE id=%s AND user_id=%s", (mid, current_user_id()))
    row = cur.fetchone()
    if not row:
        cur.close(); conn.close()
        return jsonify(message="Material not found."), 404

    cur.execute("DELETE FROM materials WHERE id=%s AND user_id=%s", (mid, current_user_id()))
    conn.commit(); cur.close(); conn.close()

    if row[0]:
        try:
            (UPLOAD_DIR / row[0]).unlink()
        except FileNotFoundError:
            pass
    return jsonify(ok=True)


@app.post("/api/reviewers")
def create_reviewer():
    if not require_user():
        return jsonify(message="Missing x-user-id header."), 401
    data = request.get_json(silent=True) or {}
    material_id = data.get("materialId")
    material = material_record(material_id, current_user_id())
    if not material:
        return jsonify(message="Material not found."), 404
    try:
        generated = generate_reviewer_from_material(material)
    except Exception as exc:
        app.logger.exception("Reviewer generation failed")
        return jsonify(message=str(exc)), 502

    title = generated["title"] or f"{Path(material['filename']).stem} Reviewer"
    reviewer = {
        "id": uid("r-"), "title": title, "sourceFile": material["filename"],
        "sourceMaterialId": material["id"], "generatedAt": now_label(),
        "subject": material["subject"],
        "mainTopics": generated.get("mainTopics", []),
        "keyConcepts": generated.get("keyConcepts", []),
        "importantTerms": generated.get("importantTerms", []),
        "summary": generated.get("summary", ""),
        "detailedContent": generated.get("detailedContent", {}),
        "isRead": False, "isBookmarked": False,
    }

    # A regeneration replaces the previous reviewer for this material, preventing
    # stale template content from remaining in the student's library.
    conn = mysql_connection(); cur = conn.cursor()
    if data.get("force"):
        cur.execute("DELETE FROM reviewers WHERE material_id=%s AND user_id=%s", (material_id, current_user_id()))
    cur.execute("""
        INSERT INTO reviewers
        (id,user_id,material_id,title,source_file,generated_at,subject,
         main_topics,key_concepts,important_terms,summary,is_read,is_bookmarked,detailed_content)
        VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
    """, (
        reviewer["id"], current_user_id(), material_id, title, material["filename"], utc_now(),
        material["subject"], jdump(reviewer["mainTopics"]), jdump(reviewer["keyConcepts"]),
        jdump(reviewer["importantTerms"]), reviewer["summary"], 0, 0, jdump(reviewer.get("detailedContent", {}))
    ))
    conn.commit(); cur.close(); conn.close()
    add_notification(current_user_id(), f"AI reviewer generated: {title}")
    return jsonify(reviewer), 201


@app.delete("/api/reviewers/<rid>")
def delete_reviewer(rid):
    if not require_user():
        return jsonify(message="Missing x-user-id header."), 401
    conn = mysql_connection(); cur = conn.cursor()
    cur.execute("DELETE FROM reviewers WHERE id=%s AND user_id=%s", (rid, current_user_id()))
    changed = cur.rowcount
    conn.commit(); cur.close(); conn.close()
    if not changed:
        return jsonify(message="Reviewer not found."), 404
    return jsonify(ok=True)


@app.patch("/api/reviewers/<rid>")
def patch_reviewer(rid):
    if not require_user():
        return jsonify(message="Missing x-user-id header."), 401
    data = request.get_json(silent=True) or {}
    allowed = {
        "isBookmarked": "is_bookmarked",
        "isRead": "is_read",
    }
    if not any(k in data for k in allowed):
        return jsonify(message="No supported reviewer fields supplied."), 400
    sets, values = [], []
    for k, col in allowed.items():
        if k in data:
            sets.append(f"{col}=%s"); values.append(int(bool(data[k])))
    values.extend([rid, current_user_id()])
    conn = mysql_connection(); cur = conn.cursor()
    cur.execute(f"UPDATE reviewers SET {', '.join(sets)} WHERE id=%s AND user_id=%s", values)
    changed = cur.rowcount
    conn.commit(); cur.close(); conn.close()
    if not changed:
        return jsonify(message="Reviewer not found."), 404
    return jsonify(next(x for x in get_user(current_user_id())["reviewers"] if x["id"] == rid))


@app.post("/api/quizzes/generate")
def generate_quiz():
    if not require_user():
        return jsonify(message="Missing x-user-id header."), 401
    data = request.get_json(silent=True) or {}
    source_type = data.get("sourceType") or "material"
    source_id = data.get("sourceId")
    requested_sources = data.get("sources") or [{"type": source_type, "id": source_id, "title": data.get("sourceTitle", "") }]
    if not isinstance(requested_sources, list) or not requested_sources or len(requested_sources) > 5:
        return jsonify(message="Select between 1 and 5 source materials."), 400
    sources = []
    for item in requested_sources:
        src = ai_source_for(item.get("type", "material"), item.get("id"), current_user_id())
        if not src:
            return jsonify(message="One of the selected source materials was not found."), 404
        sources.append(src)
    source = sources[0]
    source_titles = [str(x.get("filename") or x.get("title") or "Source") for x in sources]
    try:
        if len(sources) > 1:
            import tempfile
            combined_parts = []
            for src in sources:
                fpath = Path(src["filePath"])
                ext = fpath.suffix.lower()
                content = ""
                if ext in (".txt", ".md", ".csv", ".json", ".html", ".css", ".js", ".xml"):
                    content = fpath.read_text(encoding="utf-8", errors="replace")
                elif ext == ".pdf":
                    try:
                        from pypdf import PdfReader
                        content = "\n".join(page.extract_text() or "" for page in PdfReader(str(fpath)).pages)
                    except ImportError:
                        raise RuntimeError("Install pypdf to combine PDF sources for a quiz.")
                elif ext == ".docx":
                    from docx import Document
                    content = "\n".join(p.text for p in Document(str(fpath)).paragraphs)
                elif ext == ".pptx":
                    from pptx import Presentation
                    content = "\n".join(sh.text for slide in Presentation(str(fpath)).slides for sh in slide.shapes if getattr(sh, "has_text_frame", False))
                elif ext == ".xlsx":
                    from openpyxl import load_workbook
                    wb = load_workbook(str(fpath), read_only=True, data_only=True)
                    content = "\n".join(" ".join(str(v or "") for v in row) for ws in wb.worksheets for row in ws.iter_rows(values_only=True))
                if not content.strip():
                    raise RuntimeError(f"Could not extract readable text from {src.get('filename', 'a selected source')}. Try selecting a text-based document.")
                combined_parts.append(f"\n\n===== SOURCE: {src.get('filename','Material')} =====\n{content}")
            temp_dir = UPLOAD_DIR / "quiz-combined"
            temp_dir.mkdir(parents=True, exist_ok=True)
            combined_path = temp_dir / f"{uuid.uuid4().hex}.txt"
            combined_path.write_text("\n".join(combined_parts), encoding="utf-8")
            source = dict(source)
            source["filePath"] = str(combined_path)
            source["mimeType"] = "text/plain"
            source["filename"] = ", ".join(source_titles)
        title, questions = generate_quiz_from_material(
            source, data.get("count", 10), data.get("questionType", "mixed"),
            data.get("difficulty", "normal"), data.get("questionTypes")
        )
    except Exception as exc:
        app.logger.exception("Quiz generation failed")
        return jsonify(message=str(exc)), 502

    quiz_id = uid("quiz-")
    material_id = source["id"]
    reviewer_id = source_id if source_type == "reviewer" else None
    now = utc_now()
    conn = mysql_connection(); cur = conn.cursor()
    cur.execute("""
        INSERT INTO quizzes
        (id,user_id,material_id,reviewer_id,title,source_type,source_title,question_count,difficulty,created_at)
        VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
    """, (quiz_id, current_user_id(), material_id, reviewer_id, title,
          source_type, source["filename"], len(questions), data.get("difficulty", "normal"), now))
    for i, q in enumerate(questions, 1):
        cur.execute("""
            INSERT INTO quiz_questions
            (id,quiz_id,question_order,question_text,question_type,choices_json,correct_answer,explanation)
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s)
        """, (q["id"], quiz_id, i, q["text"], q["type"],
              jdump(q.get("choices")) if q.get("choices") else None,
              q["correctAnswer"], q.get("explanation", "")))
    conn.commit(); cur.close(); conn.close()
    add_notification(current_user_id(), f"AI quiz generated: {title}")
    return jsonify(
        quizId=quiz_id, title=title, sourceType=source_type,
        sourceTitle=", ".join(source_titles), difficulty=data.get("difficulty", "normal"), questions=questions
    ), 201


@app.post("/api/quizzes/results")
def quiz_result():
    if not require_user():
        return jsonify(message="Missing x-user-id header."), 401
    result = request.get_json(silent=True) or {}
    attempt_id = result.get("id") or uid("qa-")
    quiz_id = result.get("quizId")

    # If the frontend did not send quizId (older client), create a minimal quiz record.
    conn = mysql_connection(); cur = conn.cursor()
    if not quiz_id:
        quiz_id = uid("quiz-")
        cur.execute("""
            INSERT INTO quizzes
            (id,user_id,material_id,reviewer_id,title,source_type,source_title,question_count,difficulty,created_at)
            VALUES (%s,%s,NULL,NULL,%s,%s,%s,%s,%s,%s)
        """, (quiz_id, current_user_id(), result.get("title","Practice Quiz"),
              result.get("sourceType","material"), result.get("sourceTitle","Practice Quiz"),
              len(result.get("questions",[])), result.get("difficulty","normal"), utc_now()))
        for i, q in enumerate(result.get("questions",[]), 1):
            qid = q.get("id") or uid("q-")
            cur.execute("""
                INSERT INTO quiz_questions
                (id,quiz_id,question_order,question_text,question_type,choices_json,correct_answer,explanation)
                VALUES (%s,%s,%s,%s,%s,%s,%s,%s)
            """, (qid, quiz_id, i, q.get("text",""), q.get("type","multiple-choice"),
                  jdump(q.get("choices")) if q.get("choices") else None, q.get("correctAnswer",""), q.get("explanation")))
    completed_at = utc_now()
    cur.execute("""
        INSERT INTO quiz_attempts
        (id,user_id,quiz_id,title,source_type,source_title,total_questions,difficulty,correct_answers,score,time_spent,completed_at)
        VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)
    """, (
        attempt_id, current_user_id(), quiz_id, result.get("title","Practice Quiz"),
        result.get("sourceType","material"), result.get("sourceTitle","Practice Quiz"),
        int(result.get("totalQuestions", len(result.get("questions",[])))), result.get("difficulty","normal"),
        int(result.get("correctAnswers", 0)), float(result.get("score", 0)),
        int(result.get("timeSpent", 0)), completed_at
    ))

    answers = result.get("answers") or {}
    for q in result.get("questions", []):
        qid = q.get("id")
        selected = answers.get(qid)
        correct = q.get("correctAnswer", "")
        cur.execute("""
            INSERT INTO quiz_answers
            (id,attempt_id,question_id,question_text,selected_answer,correct_answer,is_correct)
            VALUES (%s,%s,%s,%s,%s,%s,%s)
        """, (uid("ans-"), attempt_id, qid, q.get("text",""), selected, correct,
              int(selected is not None and selected == correct)))

    conn.commit(); cur.close(); conn.close()
    add_notification(current_user_id(), f"Quiz completed: {result.get('title','Practice Quiz')}")
    result["id"] = attempt_id
    result["quizId"] = quiz_id
    result.setdefault("completedAt", now_label())
    return jsonify(result), 201


@app.patch("/api/quizzes/attempts/<attempt_id>/save")
def save_quiz_attempt(attempt_id):
    if not require_user():
        return jsonify(message="Missing x-user-id header."), 401
    saved = bool((request.get_json(silent=True) or {}).get("saved", True))
    conn = mysql_connection(); cur = conn.cursor()
    cur.execute("SELECT quiz_id FROM quiz_attempts WHERE id=%s AND user_id=%s", (attempt_id, current_user_id()))
    row = cur.fetchone()
    if not row or not row[0]:
        cur.close(); conn.close(); return jsonify(message="Quiz attempt not found."), 404
    cur.execute("UPDATE quizzes SET is_saved=%s WHERE id=%s AND user_id=%s", (int(saved), row[0], current_user_id()))
    conn.commit(); cur.close(); conn.close()
    return jsonify(id=attempt_id, isSaved=saved)


@app.patch("/api/notifications/<nid>/read")
def read_notification(nid):
    if not require_user():
        return jsonify(message="Missing x-user-id header."), 401
    conn = mysql_connection(); cur = conn.cursor()
    cur.execute("""
        UPDATE notifications SET read_flag=1
        WHERE id=%s AND user_id=%s
    """, (nid, current_user_id()))
    conn.commit(); cur.close(); conn.close()
    return jsonify(ok=True)


@app.get("/api/materials/<mid>/preview")
def preview_material(mid):
    if not require_user():
        return jsonify(message="Missing x-user-id header."), 401
    conn = mysql_connection(); cur = conn.cursor()
    cur.execute("SELECT original_filename, file_path, file_type FROM materials WHERE id=%s AND user_id=%s", (mid, current_user_id()))
    row = cur.fetchone(); cur.close(); conn.close()
    if not row:
        return jsonify(message="Material not found."), 404
    filename, file_path, file_type = row
    path = Path(file_path)
    if not path.exists(): return jsonify(message="Uploaded file is missing from storage."), 404
    ext = Path(filename).suffix.lower()
    try:
        # Images can be previewed directly by the frontend; all document formats
        # return extracted, readable text so users can inspect their real uploads.
        if request.args.get("raw") == "1":
            from flask import send_file
            return send_file(path, mimetype=mimetypes.guess_type(filename)[0] or "application/octet-stream", as_attachment=False, download_name=filename)
        content = ""
        if ext in (".txt", ".md", ".csv", ".json", ".html", ".css", ".js", ".xml"):
            content = path.read_text(encoding="utf-8", errors="replace")
        elif ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            content = "\n\n".join((page.extract_text() or "").strip() for page in reader.pages).strip()
            if not content:
                content = "No selectable text was found in this PDF. It may be a scanned document; use the PDF preview to view its pages."
        elif ext == ".docx":
            from docx import Document
            doc = Document(str(path)); content = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
            for table in doc.tables:
                content += "\n\n" + "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)
        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(str(path)); content = "\n\n".join(f"Slide {i + 1}\n" + "\n".join(sh.text for sh in slide.shapes if getattr(sh, "has_text_frame", False) and sh.text.strip()) for i, slide in enumerate(prs.slides))
        elif ext == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(str(path), read_only=True, data_only=True)
            content = "\n\n".join("Sheet: " + ws.title + "\n" + "\n".join(" | ".join(str(v or "") for v in r) for r in ws.iter_rows(values_only=True)) for ws in wb.worksheets)
        else:
            return jsonify(message="This file type does not have text preview support."), 415
        return jsonify(filename=filename, content=content, fileType=ext.lstrip("."))
    except ImportError:
        return jsonify(message="Preview dependency is not installed on the backend. Install python-docx, python-pptx, and openpyxl."), 501
    except Exception as exc:
        app.logger.exception("Document preview failed")
        return jsonify(message="Could not read this document."), 500


@app.get("/uploads/<path:filename>")
def uploads(filename):
    return send_from_directory(UPLOAD_DIR, filename)


if __name__ == "__main__":
    try:
        init_database()
        migrate_json_once()
        migrate_legacy_rows()
    except Error as e:
        print("\n[DATABASE ERROR]")
        print("Could not connect to XAMPP MySQL/MariaDB.")
        print(f"Host: {MYSQL_HOST}:{MYSQL_PORT}")
        print(f"Database: {MYSQL_DATABASE}")
        print(f"Details: {e}")
        print("\nMake sure XAMPP is running and MySQL is started.")
        raise SystemExit(1)

    print("ISLearning Python backend running at http://localhost:3001")
    print(f"MySQL/MariaDB database: {MYSQL_DATABASE} @ {MYSQL_HOST}:{MYSQL_PORT}")
    app.run(host="0.0.0.0", port=3001, debug=True)
